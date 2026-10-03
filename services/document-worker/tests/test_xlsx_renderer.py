from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.renderers.xlsx import render_xlsx


def test_fills_exact_cell_and_preserves_formula():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Wniosek"
    sheet["A1"] = "Kwota"
    sheet["B1"] = ""
    sheet["C1"] = "=1+1"

    source = BytesIO()
    workbook.save(source)

    rendered = render_xlsx(
        source.getvalue(),
        mappings=[
            {
                "fieldKey": "grant.amount",
                "locatorType": "XLSX_CELL",
                "locatorJson": {"sheet": "Wniosek", "cell": "B1"},
                "sourcePath": "Wniosek!B1",
                "inputType": "NUMBER",
            }
        ],
        values={"grant.amount": 30000},
    )

    result = load_workbook(BytesIO(rendered), data_only=False)
    assert result["Wniosek"]["A1"].value == "Kwota"
    assert result["Wniosek"]["B1"].value == 30000
    assert result["Wniosek"]["C1"].value == "=1+1"
