from io import BytesIO

from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas

from app.analyzers.docx import analyze_docx
from app.analyzers.pdf import analyze_pdf
from app.analyzers.xlsx import analyze_xlsx


def test_pdf_acroform_analysis_proposes_draft_field():
    stream = BytesIO()
    c = canvas.Canvas(stream, pagesize=(300, 300))
    c.acroForm.textfield(
        name="applicant_name",
        x=30,
        y=240,
        width=180,
        height=20,
    )
    c.save()

    result = analyze_pdf(stream.getvalue())

    assert result["analysis"]["acroFormFields"] == 1
    assert result["analysis"]["requiresManualCoordinateMapping"] is False
    assert result["mappings"][0]["locatorType"] == "PDF_ACROFORM"
    assert result["mappings"][0]["sourcePath"] == "applicant_name"


def test_docx_analysis_only_proposes_blank_cell_next_to_label():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Imię i nazwisko"
    table.cell(0, 1).text = ""
    table.cell(1, 0).text = "Stała treść"
    table.cell(1, 1).text = "Nie ruszaj"

    stream = BytesIO()
    doc.save(stream)

    result = analyze_docx(stream.getvalue())

    assert result["analysis"]["candidateFields"] == 1
    mapping = result["mappings"][0]
    assert mapping["locatorType"] == "DOCX_TABLE_CELL"
    assert mapping["questionLabel"] == "Imię i nazwisko"
    assert mapping["locatorJson"] == {"table": 0, "row": 0, "col": 1}


def test_xlsx_analysis_proposes_blank_cell_next_to_label():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Wniosek"
    sheet["A1"] = "Kwota dotacji"
    sheet["B1"] = ""
    sheet["A2"] = "Formuła"
    sheet["B2"] = "=1+1"

    stream = BytesIO()
    workbook.save(stream)

    result = analyze_xlsx(stream.getvalue())

    candidates = [
        item for item in result["mappings"]
        if item["questionLabel"] == "Kwota dotacji"
    ]
    assert len(candidates) == 1
    assert candidates[0]["sourcePath"] == "Wniosek!B1"
