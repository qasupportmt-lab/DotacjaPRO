from io import BytesIO

from docx import Document

from app.renderers.docx import render_docx


def test_fills_target_table_cell_and_preserves_other_cell():
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Nazwa działalności"
    table.cell(0, 1).text = ""
    source = BytesIO()
    doc.save(source)

    rendered = render_docx(
        source.getvalue(),
        mappings=[
            {
                "fieldKey": "business.name",
                "locatorType": "DOCX_TABLE_CELL",
                "locatorJson": {"table": 0, "row": 0, "col": 1},
                "sourcePath": "",
                "inputType": "TEXT",
            }
        ],
        values={"business.name": "Testowa Firma"},
    )

    result = Document(BytesIO(rendered))
    assert result.tables[0].cell(0, 0).text == "Nazwa działalności"
    assert result.tables[0].cell(0, 1).text == "Testowa Firma"
