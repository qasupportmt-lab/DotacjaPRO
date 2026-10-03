from io import BytesIO

from openpyxl import load_workbook

from .common import draft_key


def analyze_xlsx(source: bytes) -> dict:
    workbook = load_workbook(BytesIO(source), data_only=False)
    mappings: list[dict] = []
    order = 0

    for worksheet in workbook.worksheets:
        max_row = min(worksheet.max_row or 1, 500)
        max_col = min(worksheet.max_column or 1, 80)

        for row in range(1, max_row + 1):
            for col in range(1, max_col):
                label_cell = worksheet.cell(row=row, column=col)
                target_cell = worksheet.cell(row=row, column=col + 1)

                label = label_cell.value
                target = target_cell.value

                if not isinstance(label, str):
                    continue
                label = label.strip()
                if not label or len(label) > 300:
                    continue
                if target not in (None, ""):
                    continue

                cell_ref = target_cell.coordinate
                mappings.append({
                    "fieldKey": draft_key(
                        "xlsx",
                        label,
                        f"{worksheet.title}.{cell_ref}",
                    ),
                    "sourcePath": f"{worksheet.title}!{cell_ref}",
                    "locatorType": "XLSX_CELL",
                    "locatorJson": {
                        "sheet": worksheet.title,
                        "cell": cell_ref,
                    },
                    "inputType": "TEXT",
                    "questionLabel": label[:500],
                    "section": worksheet.title,
                    "sortOrder": order,
                    "required": False,
                    "helpText": (
                        "Pusta komórka obok etykiety wykryta automatycznie. "
                        "Wymaga kontroli przed zatwierdzeniem mapowania."
                    ),
                })
                order += 1

                if len(mappings) >= 500:
                    break
            if len(mappings) >= 500:
                break
        if len(mappings) >= 500:
            break

    return {
        "mappings": mappings,
        "analysis": {
            "format": "XLSX",
            "sheets": workbook.sheetnames,
            "candidateFields": len(mappings),
            "requiresHumanVerification": True,
            "warnings": [
                "Analiza XLSX nie zmienia formuł ani komórek; propozycje wymagają kontroli."
            ],
        },
    }
