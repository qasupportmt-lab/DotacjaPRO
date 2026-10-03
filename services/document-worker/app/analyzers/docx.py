from io import BytesIO

from docx import Document

from .common import draft_key


def _cell_text(cell) -> str:
    return " ".join(
        paragraph.text.strip()
        for paragraph in cell.paragraphs
        if paragraph.text.strip()
    ).strip()


def analyze_docx(source: bytes) -> dict:
    document = Document(BytesIO(source))
    mappings: list[dict] = []
    seen_targets: set[tuple[int, int, int]] = set()
    order = 0

    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            texts = [_cell_text(cell) for cell in row.cells]

            for col_index in range(1, len(row.cells)):
                target = texts[col_index].strip()
                label = texts[col_index - 1].strip()

                if target or not label:
                    continue
                if len(label) > 300:
                    continue

                target_id = (
                    table_index,
                    row_index,
                    col_index,
                )
                if target_id in seen_targets:
                    continue
                seen_targets.add(target_id)

                mappings.append({
                    "fieldKey": draft_key(
                        "docx",
                        label,
                        f"{table_index}.{row_index}.{col_index}",
                    ),
                    "sourcePath": (
                        f"table:{table_index}/row:{row_index}/col:{col_index}"
                    ),
                    "locatorType": "DOCX_TABLE_CELL",
                    "locatorJson": {
                        "table": table_index,
                        "row": row_index,
                        "col": col_index,
                    },
                    "inputType": "TEXT",
                    "questionLabel": label[:500],
                    "section": f"Tabela {table_index + 1}",
                    "sortOrder": order,
                    "required": False,
                    "helpText": (
                        "Pusta komórka obok etykiety wykryta automatycznie. "
                        "Wymaga kontroli przed zatwierdzeniem mapowania."
                    ),
                })
                order += 1

    return {
        "mappings": mappings,
        "analysis": {
            "format": "DOCX",
            "paragraphs": len(document.paragraphs),
            "tables": len(document.tables),
            "candidateFields": len(mappings),
            "requiresHumanVerification": True,
            "warnings": [
                "Mapowanie komórek DOCX jest wyłącznie propozycją techniczną."
            ],
        },
    }
