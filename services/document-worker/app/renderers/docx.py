from io import BytesIO
from typing import Any

from docx import Document

from ..values import as_text


def _set_paragraph_text(paragraph, value: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


def _replace_token_in_paragraph(paragraph, token: str, value: str) -> bool:
    full_text = "".join(run.text for run in paragraph.runs) if paragraph.runs else paragraph.text
    if token not in full_text:
        return False
    _set_paragraph_text(paragraph, full_text.replace(token, value))
    return True


def _iter_all_paragraphs(document):
    for paragraph in document.paragraphs:
        yield paragraph

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def render_docx(
    source: bytes,
    mappings: list[dict],
    values: dict[str, Any],
) -> bytes:
    document = Document(BytesIO(source))

    for mapping in mappings:
        field_key = mapping["fieldKey"]
        if field_key not in values:
            continue

        value = as_text(values[field_key], mapping.get("inputType"))
        locator_type = mapping["locatorType"]
        locator = mapping.get("locatorJson") or {}

        if locator_type == "DOCX_TOKEN":
            token = mapping.get("sourcePath")
            if not token:
                raise ValueError(f"Missing DOCX token for {field_key}")

            replaced = False
            for paragraph in _iter_all_paragraphs(document):
                replaced = _replace_token_in_paragraph(paragraph, token, value) or replaced

            if not replaced:
                raise ValueError(f"DOCX token not found: {token}")

        elif locator_type == "DOCX_TABLE_CELL":
            table_index = int(locator["table"])
            row_index = int(locator["row"])
            col_index = int(locator["col"])

            try:
                cell = document.tables[table_index].rows[row_index].cells[col_index]
            except (IndexError, KeyError) as exc:
                raise ValueError(f"Invalid DOCX table locator for {field_key}") from exc

            if cell.paragraphs:
                _set_paragraph_text(cell.paragraphs[0], value)
                for paragraph in cell.paragraphs[1:]:
                    _set_paragraph_text(paragraph, "")
            else:
                cell.add_paragraph(value)

        elif locator_type == "DOCX_PARAGRAPH":
            paragraph_index = int(locator["paragraph"])
            try:
                paragraph = document.paragraphs[paragraph_index]
            except (IndexError, KeyError) as exc:
                raise ValueError(f"Invalid DOCX paragraph locator for {field_key}") from exc
            _set_paragraph_text(paragraph, value)

    output = BytesIO()
    document.save(output)
    return output.getvalue()
