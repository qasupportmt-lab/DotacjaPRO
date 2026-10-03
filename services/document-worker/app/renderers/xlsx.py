from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from ..values import as_text


def render_xlsx(
    source: bytes,
    mappings: list[dict],
    values: dict[str, Any],
) -> bytes:
    workbook = load_workbook(BytesIO(source))

    for mapping in mappings:
        if mapping["locatorType"] != "XLSX_CELL":
            continue

        field_key = mapping["fieldKey"]
        if field_key not in values:
            continue

        locator = mapping.get("locatorJson") or {}
        sheet_name = locator.get("sheet")
        cell_ref = locator.get("cell")

        if not sheet_name or not cell_ref:
            source_path = mapping.get("sourcePath", "")
            if "!" not in source_path:
                raise ValueError(f"Invalid XLSX locator for {field_key}")
            sheet_name, cell_ref = source_path.rsplit("!", 1)

        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Sheet not found: {sheet_name}")

        worksheet = workbook[sheet_name]
        value = values[field_key]

        if isinstance(value, (int, float, bool)):
            worksheet[cell_ref] = value
        else:
            worksheet[cell_ref] = as_text(value, mapping.get("inputType"))

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
